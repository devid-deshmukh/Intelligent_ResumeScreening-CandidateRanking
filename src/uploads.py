"""Extract and normalize uploaded resumes for the existing ranking engine."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

from src.preprocess import clean_text


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from one supported PDF or DOCX file."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF and DOCX files are supported.")
    if not content:
        raise ValueError("The file is empty.")
    try:
        if extension == ".pdf":
            try:
                from pypdf import PdfReader
            except ModuleNotFoundError as error:
                raise ValueError(
                    "PDF support is not installed in the active Python environment. "
                    "Run venv\\Scripts\\python.exe -m pip install -r requirements.txt, then restart Streamlit."
                ) from error

            pages = PdfReader(io.BytesIO(content)).pages
            text = "\n".join(page.extract_text() or "" for page in pages)
        else:
            try:
                from docx import Document
            except ModuleNotFoundError as error:
                raise ValueError(
                    "DOCX support is not installed in the active Python environment. "
                    "Run venv\\Scripts\\python.exe -m pip install -r requirements.txt, then restart Streamlit."
                ) from error

            document = Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = [" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
            text = "\n".join(paragraphs + tables)
    except Exception as error:
        raise ValueError(f"Could not read the {extension[1:].upper()} file: {error}") from error
    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the file.")
    return text


def _skill_matches(text: str, skills: Iterable[str]) -> list[str]:
    normalized = clean_text(text)
    matches = []
    for skill in skills:
        skill_text = clean_text(skill)
        if skill_text and re.search(rf"\b{re.escape(skill_text)}\b", normalized):
            matches.append(skill_text)
    return list(dict.fromkeys(matches))


def _years_experience(text: str) -> float:
    patterns = (
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+(?:of\s+)?experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*years?",
    )
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return float(match.group(1))
    return 0.0


def _seniority(text: str) -> str:
    normalized = clean_text(text)
    for value in ("principal", "lead", "senior", "mid level", "junior", "entry level"):
        if re.search(rf"\b{re.escape(value)}\b", normalized):
            return value
    return ""


def resume_record(filename: str, text: str, job: dict[str, Any], resume_id: str) -> dict[str, Any]:
    """Convert extracted text to the columns consumed by ``rank_resumes``."""
    required = job.get("must_have_skills", [])
    preferred = job.get("nice_to_have_skills", [])
    return {
        "resume_id": resume_id,
        "filename": filename,
        "role": text.splitlines()[0].strip()[:120] if text.splitlines() else "",
        "seniority": _seniority(text),
        "industry": "",
        "years_experience": _years_experience(text),
        "skills": _skill_matches(text, list(required) + list(preferred)),
        "summary": text,
        "resume_text": text,
    }


def process_uploaded_resumes(files: Iterable[Any], job: dict[str, Any]) -> tuple[pd.DataFrame, list[dict[str, str]]]:
    """Process all files, returning valid records and per-file errors."""
    records: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    seen_ids: dict[str, int] = {}
    for uploaded_file in files:
        filename = str(getattr(uploaded_file, "name", "uploaded_resume"))
        stem = Path(filename).stem or "resume"
        count = seen_ids.get(stem.casefold(), 0) + 1
        seen_ids[stem.casefold()] = count
        resume_id = stem if count == 1 else f"{stem}_{count}"
        try:
            content = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else uploaded_file.read()
            text = extract_resume_text(filename, content)
            records.append(resume_record(filename, text, job, resume_id))
        except (OSError, ValueError) as error:
            errors.append({"filename": filename, "error": str(error)})
    return pd.DataFrame(records), errors
